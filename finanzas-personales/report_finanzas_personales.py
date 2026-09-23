# report_finanzas_personales.py
# Exporter (TXT + PDF) para la Consultoría de Finanzas Personales — sin matplotlib, PDF 100% ReportLab

from __future__ import annotations

import io
import math
import textwrap
from dataclasses import dataclass, field
from typing import Optional, Sequence, Tuple

REPORTLAB_OK = True
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas
except Exception:
    REPORTLAB_OK = False

DOCX_OK = True
try:
    import docx
    from docx.shared import Pt, RGBColor
except Exception:
    DOCX_OK = False


# ----------------------------
# Formatos
# ----------------------------
def fmt_pct(x: Optional[float]) -> str:
    if x is None or not math.isfinite(x):
        return "—"
    return f"{x*100:.1f}%"

def fmt_pyg(x: Optional[float]) -> str:
    if x is None or not math.isfinite(x):
        return "—"
    return "Gs. {:,.0f}".format(float(x)).replace(",", ".")

def wrap(s: str, width: int) -> list[str]:
    return textwrap.wrap(s, width=width, break_long_words=False, replace_whitespace=False)


# ----------------------------
# Data model
# ----------------------------
@dataclass
class Rating:
    key: str      # "good" | "warn" | "bad" | "na"
    label: str
    emoji: str

@dataclass
class DebtItem:
    nombre: str
    saldo: float
    tasa_anual: float   # ej. 0.24 = 24% anual
    pago_minimo: float

@dataclass
class ClientFinReport:
    consultant: str
    client: str
    report_date: str
    age: Optional[int]
    occupation: str
    dependents: int
    objective: str

    # Perfil de riesgo
    risk_score: int
    risk_max: int
    risk_category: str
    risk_description: str
    risk_alloc_fixed: int   # % ilustrativo en instrumentos de bajo riesgo
    risk_alloc_variable: int  # % ilustrativo en instrumentos de mayor riesgo

    # Diagnóstico financiero (montos)
    income: float
    fixed_expenses: float
    variable_expenses: float
    debt_payment: float
    savings_monthly: float
    emergency_fund: float
    net_worth: Optional[float]

    # Indicadores calculados
    total_expenses: float
    cashflow: float
    savings_rate: float
    dti: float
    emergency_months: float

    cashflow_rating: Rating
    savings_rating: Rating
    dti_rating: Rating
    emergency_rating: Rating

    health_score: int
    health_max: int
    health_label: str

    action_plan: Sequence[str] = field(default_factory=list)

    # --- Secciones opcionales (todas vacías/None por defecto => no aparecen en el
    # reporte ni afectan a nadie que no las use; ver página 2 del PDF / TXT). ---
    # Meta de ahorro
    savings_goal_name: Optional[str] = None
    savings_goal_amount: Optional[float] = None
    savings_goal_months: Optional[int] = None
    savings_goal_required_monthly: Optional[float] = None
    savings_goal_feasible: Optional[bool] = None

    # Inventario de deudas + priorización
    debts: Sequence[DebtItem] = field(default_factory=list)
    debt_priority_note: Optional[str] = None

    # Plan editable: compromisos acordados (texto, fecha, responsable)
    commitments: Sequence[Tuple[str, str, str]] = field(default_factory=list)

    # Cartera modelo sugerida (por categoría de instrumento, no títulos específicos)
    portfolio_tier: Optional[str] = None  # "Conservadora" | "Moderada" | "Arriesgada"
    portfolio_categories: Sequence[Tuple[str, int]] = field(default_factory=list)
    portfolio_age_note: Optional[str] = None

DISCLAIMER = (
    "Herramienta educativa del Diplomado de Finanzas Personales. No constituye asesoría de inversión "
    "regulada; valida cualquier decisión relevante con un asesor financiero certificado."
)

# El cuestionario de 10 preguntas mide sobre todo TOLERANCIA al riesgo (actitud/disposición
# psicológica frente a la volatilidad). No sustituye un análisis de CAPACIDAD de riesgo (qué
# tanta pérdida puede absorber el cliente según edad, estabilidad de ingresos, dependientes,
# horizonte y fondo de emergencia) — ambos conceptos son distintos en planificación financiera
# y deberían combinarse antes de sugerir una mezcla de inversión. Ver recomendación al docente.
RISK_METHOD_NOTE = (
    "Nota metodológica: este puntaje refleja principalmente la tolerancia al riesgo (actitud) del "
    "cliente. Antes de definir una mezcla de inversión, contrasta este resultado con su capacidad de "
    "riesgo real (edad, estabilidad de ingresos, dependientes y fondo de emergencia)."
)

PORTFOLIO_DISCLAIMER = (
    "Cartera modelo con fines educativos, por categoría de instrumento financiero (no recomienda "
    "títulos, emisores ni productos específicos). No es asesoría de inversión personalizada ni "
    "sustituye la debida diligencia de un asesor financiero certificado."
)


# ----------------------------
# TXT (para descarga rápida)
# ----------------------------
def build_text_report(r: ClientFinReport) -> str:
    lines: list[str] = []
    lines.append("DIAGNÓSTICO DE FINANZAS PERSONALES — REPORTE DE CONSULTORÍA")
    lines.append(f"Consultor/a: {r.consultant}  |  Cliente: {r.client}")
    lines.append(f"Fecha: {r.report_date}")
    edad = "—" if r.age is None else str(r.age)
    lines.append(f"Edad: {edad}  |  Ocupación: {r.occupation}  |  Dependientes: {r.dependents}")
    lines.append(f"Objetivo principal: {r.objective}")
    lines.append("")
    lines.append(f"PERFIL DE RIESGO: {r.risk_category}  ({r.risk_score}/{r.risk_max} pts)")
    lines.append(r.risk_description)
    lines.append(
        f"Mezcla ilustrativa (no es recomendación de inversión): "
        f"{r.risk_alloc_fixed}% bajo riesgo/ahorro — {r.risk_alloc_variable}% mayor riesgo/crecimiento"
    )
    lines.append(RISK_METHOD_NOTE)

    if r.portfolio_tier and r.portfolio_categories:
        lines.append("")
        lines.append(f"CARTERA MODELO SUGERIDA: {r.portfolio_tier}")
        if r.portfolio_age_note:
            lines.append(f"({r.portfolio_age_note})")
        for categoria, pct in r.portfolio_categories:
            lines.append(f"- {categoria}: {pct}%")
        lines.append(PORTFOLIO_DISCLAIMER)

    lines.append("")
    lines.append("DIAGNÓSTICO FINANCIERO RÁPIDO")
    lines.append(f"- Ingreso mensual: {fmt_pyg(r.income)}")
    lines.append(f"- Gastos fijos: {fmt_pyg(r.fixed_expenses)} | Gastos variables: {fmt_pyg(r.variable_expenses)}")
    lines.append(f"- Cuota de deudas: {fmt_pyg(r.debt_payment)}")
    lines.append(f"- Ahorro/inversión mensual: {fmt_pyg(r.savings_monthly)}")
    lines.append(f"- Fondo de emergencia actual: {fmt_pyg(r.emergency_fund)}")
    if r.net_worth is not None:
        lines.append(f"- Patrimonio neto aproximado: {fmt_pyg(r.net_worth)}")
    lines.append("")
    lines.append("Indicadores")
    lines.append(f"- Excedente antes de ahorro: {fmt_pyg(r.cashflow)}  [{r.cashflow_rating.emoji} {r.cashflow_rating.label}]")
    savings_txt = "No calculable (sin ingreso)" if r.savings_rating.key == "na" else fmt_pct(r.savings_rate)
    lines.append(f"- Tasa de ahorro: {savings_txt}  [{r.savings_rating.emoji} {r.savings_rating.label}]")
    dti_txt = "No calculable (sin ingreso)" if r.dti_rating.key == "na" else fmt_pct(r.dti)
    lines.append(f"- Endeudamiento (cuota/ingreso): {dti_txt}  [{r.dti_rating.emoji} {r.dti_rating.label}]")
    lines.append(f"- Fondo de emergencia: {r.emergency_months:.1f} meses cubiertos  [{r.emergency_rating.emoji} {r.emergency_rating.label}]")
    available_after_savings = r.cashflow - r.savings_monthly
    lines.append(f"- Disponible después de ahorro (excedente − ahorro declarado): {fmt_pyg(available_after_savings)}")
    if available_after_savings < 0:
        lines.append(
            "  ⚠️ El ahorro declarado supera el excedente disponible; verificar con el cliente el origen de "
            "esos fondos (otros ingresos, activos existentes o deuda adicional)."
        )
    lines.append("")
    lines.append(f"SALUD FINANCIERA GENERAL: {r.health_label} ({r.health_score}/{r.health_max})")
    lines.append("")
    lines.append("PLAN DE ACCIÓN SUGERIDO")
    for item in r.action_plan:
        lines.append(f"- {item}")

    if r.savings_goal_amount is not None and r.savings_goal_months:
        lines.append("")
        lines.append("META DE AHORRO")
        nombre_meta = r.savings_goal_name or "Meta sin nombre"
        lines.append(f"- {nombre_meta}: {fmt_pyg(r.savings_goal_amount)} en {r.savings_goal_months} meses")
        lines.append(f"- Ahorro mensual necesario: {fmt_pyg(r.savings_goal_required_monthly)}")
        lines.append(f"- Ahorro mensual actual declarado: {fmt_pyg(r.savings_monthly)}")
        veredicto = "Alcanzable con el ahorro actual" if r.savings_goal_feasible else "Requiere aumentar el ahorro mensual o extender el plazo"
        lines.append(f"- Veredicto: {veredicto}")

    if r.debts:
        lines.append("")
        lines.append("INVENTARIO DE DEUDAS")
        for d in r.debts:
            lines.append(
                f"- {d.nombre}: saldo {fmt_pyg(d.saldo)} | tasa {fmt_pct(d.tasa_anual)} anual | "
                f"pago mínimo {fmt_pyg(d.pago_minimo)}"
            )
        if r.debt_priority_note:
            lines.append(f"- {r.debt_priority_note}")

    if r.commitments:
        lines.append("")
        lines.append("COMPROMISOS ACORDADOS")
        for texto, fecha, responsable in r.commitments:
            lines.append(f"- {texto} | Fecha: {fecha or '—'} | Responsable: {responsable or '—'}")

    lines.append("")
    lines.append(DISCLAIMER)
    return "\n".join(lines)


# ----------------------------
# PDF (ReportLab, estilo dark dashboard)
# ----------------------------
def generate_pdf(r: ClientFinReport) -> bytes:
    if not REPORTLAB_OK:
        raise RuntimeError("ReportLab no está disponible. Agregar `reportlab` a requirements.txt.")

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    W, H = letter

    bg = colors.HexColor("#050914")
    card = colors.HexColor("#0b1733")
    card2 = colors.HexColor("#0d1b3d")
    line = colors.Color(1, 1, 1, alpha=0.10)
    text = colors.HexColor("#EAF1FF")
    muted = colors.Color(234/255, 241/255, 1, alpha=0.75)
    accent = colors.HexColor("#66A9FF")
    good = colors.HexColor("#27D17C")
    warn = colors.HexColor("#FFCC66")
    bad = colors.HexColor("#FF5D5D")

    rating_color = {"good": good, "warn": warn, "bad": bad}

    def rr(x, y, w, h, r=14, fill=card):
        c.setFillColor(fill)
        c.setStrokeColor(line)
        c.setLineWidth(1)
        c.roundRect(x, y, w, h, r, stroke=1, fill=1)

    def t(x, y, s, size=10.5, bold=False, col=text):
        c.setFillColor(col)
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(x, y, s)

    def tr(x, y, s, size=10.5, bold=False, col=text):
        c.setFillColor(col)
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawRightString(x, y, s)

    def para(x, y, s, width_chars=90, size=9.4, leading=12, col=muted):
        yy = y
        for ln in wrap(s, width_chars):
            t(x, yy, ln, size=size, bold=False, col=col)
            yy -= leading
        return yy

    def ellipsize(s, max_width, size=9.5, bold=False):
        # Trunca con "…" el texto que no entra en max_width (puntos), en vez de
        # dejar que invada lo que está dibujado a la derecha (fecha, píldora).
        font = "Helvetica-Bold" if bold else "Helvetica"
        if c.stringWidth(s, font, size) <= max_width:
            return s
        ell = "…"
        trimmed = s
        while trimmed and c.stringWidth(trimmed + ell, font, size) > max_width:
            trimmed = trimmed[:-1]
        return (trimmed.rstrip() + ell) if trimmed else ell

    # Background
    c.setFillColor(bg)
    c.rect(0, 0, W, H, stroke=0, fill=1)

    margin = 0.55 * inch
    top = H - margin
    left = margin
    right = W - margin

    # Header
    # Altura ampliada (62 -> 80) para poder mostrar el objetivo del cliente en una
    # línea propia; kpi_y se corre la misma cantidad (18pt) más abajo para que el
    # espaciado con la fila de KPIs no cambie.
    header_h = 80
    rr(left, top-88, right-left, header_h, r=16, fill=card2)
    t(left+16, top-28, "DIAGNÓSTICO DE FINANZAS PERSONALES — REPORTE DE CONSULTORÍA", size=12.5, bold=True)
    edad = "—" if r.age is None else str(r.age)

    # La píldora de salud financiera (dibujada más abajo) ocupa una franja vertical
    # que cruza tanto la fila 1 (Consultor/a | Cliente) como la fila 2 (Edad |
    # Ocupación | Dependientes), así que ambas deben frenar antes de su borde
    # izquierdo. La fecha, en cambio, NO va en esas filas: si se dibuja a la misma
    # altura que la píldora (right-63 a top-37), la píldora se pinta encima y la
    # tapa (se dibuja después). Por eso la fecha se movió a la fila del objetivo,
    # que queda por debajo del borde inferior de la píldora.
    pill_left_x = right-210
    date_str = f"Fecha: {r.report_date}"
    date_w = c.stringWidth(date_str, "Helvetica", 9.5)
    date_start_x = right-16 - date_w

    line1_max_w = pill_left_x - 10 - (left+16)
    line1 = ellipsize(f"Consultor/a: {r.consultant}  |  Cliente: {r.client}", line1_max_w, size=9.5)
    t(left+16, top-45, line1, size=9.5, col=muted)

    line2_max_w = pill_left_x - 10 - (left+16)
    line2 = ellipsize(
        f"Edad: {edad}  |  Ocupación: {r.occupation}  |  Dependientes: {r.dependents}", line2_max_w, size=9.5,
    )
    t(left+16, top-60, line2, size=9.5, col=muted)

    line3_max_w = date_start_x - 12 - (left+16)
    line3 = ellipsize(f"Objetivo: {r.objective}", line3_max_w, size=9.5)
    t(left+16, top-75, line3, size=9.5, col=accent)
    tr(right-16, top-75, date_str, size=9.5, col=muted)

    # Health pill
    pill_col = good if r.health_score >= 6 else (bad if r.health_score <= 3 else warn)
    c.setFillColor(colors.Color(pill_col.red, pill_col.green, pill_col.blue, alpha=0.18))
    c.setStrokeColor(line)
    c.roundRect(right-210, top-63, 190, 26, 13, stroke=1, fill=1)
    t(right-200, top-56, f"SALUD FINANCIERA: {r.health_label}", size=9.6, bold=True, col=pill_col)

    # KPI row (4 cards)
    kpi_y = top-168
    kpi_h = 58
    gap = 10
    kpi_w = (right-left - gap*3)/4

    def kpi(i, title, value, rating: Rating):
        x = left + i*(kpi_w+gap)
        rr(x, kpi_y, kpi_w, kpi_h, r=14, fill=card)
        t(x+12, kpi_y+kpi_h-18, title, size=8.4, col=muted)
        t(x+12, kpi_y+20, value, size=11.5, bold=True)
        t(x+12, kpi_y+6, f"{rating.emoji} {rating.label}", size=8.4, col=rating_color.get(rating.key, muted))

    savings_val = "—" if r.savings_rating.key == "na" else fmt_pct(r.savings_rate)
    dti_val = "—" if r.dti_rating.key == "na" else fmt_pct(r.dti)
    emergency_val = "—" if r.emergency_rating.key == "na" else f"{r.emergency_months:.1f} meses"
    kpi(0, "Excedente antes de ahorro", fmt_pyg(r.cashflow), r.cashflow_rating)
    kpi(1, "Tasa de ahorro", savings_val, r.savings_rating)
    kpi(2, "Endeudamiento (DTI)", dti_val, r.dti_rating)
    kpi(3, "Fondo de emergencia", emergency_val, r.emergency_rating)

    # Mid cards: perfil de riesgo (izq) + diagnóstico (der)
    # Altura dinámica: antes la tarjeta tenía una altura fija (190pt) que asumía una
    # descripción de riesgo corta; en la práctica CASI TODAS las categorías (4 de 5,
    # más el estado "Pendiente") envuelven a 4 líneas y, sumadas a la nota
    # metodológica (5 líneas), se salían de esa altura fija — la barra de "mezcla
    # ilustrativa" terminaba dibujada encima del texto. Ahora la tarjeta mide el
    # contenido real (líneas de texto envueltas) y crece hacia abajo si hace falta,
    # empujando el resto del layout (que ya depende de mid_y) sin solaparse.
    desc_lines = wrap(r.risk_description, 52)
    note_lines = wrap(RISK_METHOD_NOTE, 60)
    caption_text = (
        f"Mezcla ilustrativa: {r.risk_alloc_fixed}% bajo riesgo / {r.risk_alloc_variable}% mayor riesgo "
        f"(no es recomendación)"
    )
    caption_leading = 9.2
    caption_lines = wrap(caption_text, 58)
    left_needed_h = (
        58 + len(desc_lines) * 12 + 4 + len(note_lines) * 9.5 + 12 + 14 + 12
        + len(caption_lines) * caption_leading + 14
    )
    right_rows_count = 7 if r.net_worth is not None else 6
    right_needed_h = 44 + right_rows_count * 15 + 10

    mid_h = max(190, left_needed_h, right_needed_h)
    mid_top = kpi_y - 20
    mid_y = mid_top - mid_h
    lw = (right-left)*0.50-8
    rw = (right-left)*0.50-8
    rr(left, mid_y, lw, mid_h, r=18, fill=card)
    rr(left+lw+16, mid_y, rw, mid_h, r=18, fill=card)

    # Izquierda: perfil de riesgo (todo posicionado desde el borde superior fijo
    # mid_top, no desde mid_y+mid_h, para que el contenido no se mueva aunque la
    # tarjeta crezca hacia abajo)
    t(left+16, mid_top-24, "Perfil de riesgo", size=11.2, bold=True)
    t(left+16, mid_top-40, f"{r.risk_category}  —  {r.risk_score}/{r.risk_max} pts", size=10, bold=True, col=accent)
    yy = para(left+16, mid_top-58, r.risk_description, width_chars=52)
    yy = para(left+16, yy - 4, RISK_METHOD_NOTE, width_chars=60, size=7.6, leading=9.5, col=accent)

    bar_x = left+16
    bar_h = 14
    # Ya no hace falta acotar la posición con max()/min(): la tarjeta mide lo que el
    # texto necesita, así que la barra simplemente va justo debajo de la nota.
    bar_y = yy - 12
    bar_w = lw-32
    c.setFillColor(colors.Color(1, 1, 1, alpha=0.05))
    c.rect(bar_x, bar_y, bar_w, bar_h, stroke=0, fill=1)
    seg1 = bar_w * (r.risk_alloc_fixed/100.0)
    c.setFillColor(colors.Color(accent.red, accent.green, accent.blue, alpha=0.55))
    c.rect(bar_x, bar_y, seg1, bar_h, stroke=0, fill=1)
    c.setFillColor(colors.Color(warn.red, warn.green, warn.blue, alpha=0.55))
    c.rect(bar_x+seg1, bar_y, bar_w-seg1, bar_h, stroke=0, fill=1)
    cap_y = bar_y - 12
    for ln in caption_lines:
        t(left+16, cap_y, ln, size=7.6, col=muted)
        cap_y -= caption_leading

    # Derecha: diagnóstico financiero
    rx = left+lw+16
    t(rx+16, mid_top-24, "Diagnóstico financiero", size=11.2, bold=True)
    rows = [
        ("Ingreso mensual", fmt_pyg(r.income)),
        ("Gastos fijos", fmt_pyg(r.fixed_expenses)),
        ("Gastos variables", fmt_pyg(r.variable_expenses)),
        ("Cuota de deudas", fmt_pyg(r.debt_payment)),
        ("Ahorro/inversión mensual", fmt_pyg(r.savings_monthly)),
        ("Fondo de emergencia (monto)", fmt_pyg(r.emergency_fund)),
    ]
    if r.net_worth is not None:
        rows.append(("Patrimonio neto aprox.", fmt_pyg(r.net_worth)))
    available_after_savings = r.cashflow - r.savings_monthly
    rows.append(("Disponible después de ahorro", fmt_pyg(available_after_savings)))
    ry = mid_top-44
    for label, val in rows:
        t(rx+16, ry, label, size=8.9, col=muted)
        val_col = bad if (label == "Disponible después de ahorro" and available_after_savings < 0) else text
        tr(rx+rw-16, ry, val, size=8.9, bold=True, col=val_col)
        ry -= 15

    # Bottom card: plan de acción
    # Usa todo el espacio vertical libre entre las tarjetas del medio y el pie de
    # página (antes quedaba una franja en blanco y una tarjeta de altura fija que
    # truncaba planes de acción largos sin avisar).
    footer_reserved = 34  # línea de marca + disclaimer (ver footer más abajo)
    bot_top = mid_y - 20
    bot_y = margin + footer_reserved
    bot_h = bot_top - bot_y
    rr(left, bot_y, right-left, bot_h, r=18, fill=card)
    t(left+16, bot_y+bot_h-24, "Plan de acción sugerido", size=11.2, bold=True)
    yy = bot_y+bot_h-44
    truncated = False
    items = list(r.action_plan)
    for idx, item in enumerate(items):
        wrapped_lines = wrap(item, 108)
        stop = False
        for line_idx, ln in enumerate(wrapped_lines):
            # Antes de dibujar cada línea (no recién al terminar el ítem) chequeamos si
            # queda contenido pendiente y ya no hay lugar: si dibujáramos igual, la última
            # línea del ítem podía terminar superpuesta con el aviso de continuación.
            more_content = (line_idx < len(wrapped_lines) - 1) or (idx < len(items) - 1)
            if yy < bot_y + 24 and more_content:
                truncated = True
                stop = True
                break
            t(left+20, yy, ln, size=9.2, col=muted)
            yy -= 12.5
        if stop:
            break
        yy -= 2
    if truncated:
        t(left+20, bot_y+12, "(continúa en el reporte TXT — ver ítems adicionales)", size=8.2, col=accent)

    def draw_footer():
        c.setFillColor(muted)
        c.setFont("Helvetica", 8.2)
        disclaimer_lines = wrap(DISCLAIMER, 100)
        brand_y = margin - 8 + (len(disclaimer_lines) - 1) * 10
        c.drawRightString(right, brand_y, "Consultoría de Finanzas Personales — Diplomado")
        for i, ln in enumerate(disclaimer_lines):
            c.drawString(left, margin-8-(i*10), ln)

    draw_footer()

    # --------------------------------------------------------------
    # Página 2: se agrega si hay cartera modelo, meta de ahorro, deudas o
    # compromisos cargados. La cartera modelo sugerida se calcula siempre que
    # el test de riesgo esté completo, así que en la práctica el PDF pasa a
    # ser de 2 páginas en el caso normal (antes era de 1 sola cuando nadie
    # usaba las secciones opcionales).
    # --------------------------------------------------------------
    has_page2 = (
        bool(r.portfolio_tier and r.portfolio_categories)
        or bool(r.savings_goal_amount is not None and r.savings_goal_months)
        or bool(r.debts) or bool(r.commitments)
    )
    if has_page2:
        c.showPage()
        c.setFillColor(bg)
        c.rect(0, 0, W, H, stroke=0, fill=1)

        y = top
        t(left, y, "DIAGNÓSTICO DE FINANZAS PERSONALES — PÁGINA 2", size=11.5, bold=True)
        tr(right, y, f"{r.consultant}  |  {r.client}", size=8.8, col=muted)
        y -= 26

        def section_card(title, min_h, draw_body):
            nonlocal y
            top_y = y
            body_bottom = draw_body(top_y - 34)
            h = max(min_h, (top_y - body_bottom) + 14)
            rr(left, top_y - h, right-left, h, r=16, fill=card)
            t(left+16, top_y-24, title, size=11.2, bold=True)
            draw_body(top_y - 34)
            y = top_y - h - 16

        if r.portfolio_tier and r.portfolio_categories:
            def body_portfolio(yy0):
                yy = yy0
                t(left+16, yy, f"Nivel sugerido: {r.portfolio_tier}", size=9.6, bold=True, col=accent)
                yy -= 16
                if r.portfolio_age_note:
                    yy = para(left+16, yy, r.portfolio_age_note, width_chars=112, size=8.4, leading=10.5, col=muted)
                    yy -= 4

                palette = [
                    colors.Color(good.red, good.green, good.blue, alpha=0.55),
                    colors.Color(accent.red, accent.green, accent.blue, alpha=0.75),
                    colors.Color(accent.red, accent.green, accent.blue, alpha=0.40),
                    colors.Color(warn.red, warn.green, warn.blue, alpha=0.55),
                    colors.Color(bad.red, bad.green, bad.blue, alpha=0.55),
                ]
                bar_x = left+16
                bar_w = (right-left) - 32
                bar_h = 16
                c.setFillColor(colors.Color(1, 1, 1, alpha=0.05))
                c.rect(bar_x, yy-bar_h, bar_w, bar_h, stroke=0, fill=1)
                x_cursor = bar_x
                for i, (_categoria, pct) in enumerate(r.portfolio_categories):
                    seg_w = bar_w * (pct/100.0)
                    c.setFillColor(palette[i % len(palette)])
                    c.rect(x_cursor, yy-bar_h, seg_w, bar_h, stroke=0, fill=1)
                    x_cursor += seg_w
                yy -= bar_h + 14

                for i, (categoria, pct) in enumerate(r.portfolio_categories):
                    c.setFillColor(palette[i % len(palette)])
                    c.rect(left+16, yy-8, 10, 10, stroke=0, fill=1)
                    t(left+32, yy, f"{categoria}: {pct}%", size=9.2, col=text)
                    yy -= 15
                yy -= 4
                yy = para(left+16, yy, PORTFOLIO_DISCLAIMER, width_chars=115, size=7.8, leading=10, col=muted)
                return yy
            section_card("Cartera modelo sugerida (por categoría de instrumento)", 150, body_portfolio)

        if r.savings_goal_amount is not None and r.savings_goal_months:
            def body_goal(yy0):
                nombre_meta = r.savings_goal_name or "Meta sin nombre"
                yy = yy0
                t(left+16, yy, f"{nombre_meta}: {fmt_pyg(r.savings_goal_amount)} en {r.savings_goal_months} meses", size=9.6, bold=True)
                yy -= 16
                t(left+16, yy, f"Ahorro mensual necesario: {fmt_pyg(r.savings_goal_required_monthly)}", size=9.2, col=muted)
                yy -= 14
                t(left+16, yy, f"Ahorro mensual actual declarado: {fmt_pyg(r.savings_monthly)}", size=9.2, col=muted)
                yy -= 14
                veredicto = "✅ Alcanzable con el ahorro actual" if r.savings_goal_feasible else "⚠️ Requiere aumentar el ahorro o extender el plazo"
                vcol = good if r.savings_goal_feasible else warn
                t(left+16, yy, veredicto, size=9.4, bold=True, col=vcol)
                return yy - 10
            section_card("Meta de ahorro", 90, body_goal)

        if r.debts:
            def body_debts(yy0):
                yy = yy0
                t(left+16, yy, "Deuda", size=8.6, col=muted, bold=True)
                t(left+220, yy, "Saldo", size=8.6, col=muted, bold=True)
                t(left+340, yy, "Tasa anual", size=8.6, col=muted, bold=True)
                tr(right-16, yy, "Pago mínimo", size=8.6, col=muted, bold=True)
                yy -= 14
                for d in r.debts:
                    t(left+16, yy, ellipsize(d.nombre, 195, size=9.0), size=9.0)
                    t(left+220, yy, fmt_pyg(d.saldo), size=9.0)
                    t(left+340, yy, fmt_pct(d.tasa_anual), size=9.0)
                    tr(right-16, yy, fmt_pyg(d.pago_minimo), size=9.0)
                    yy -= 14
                if r.debt_priority_note:
                    yy -= 4
                    yy = para(left+16, yy, r.debt_priority_note, width_chars=100, size=8.8, leading=11, col=accent)
                return yy
            section_card("Inventario de deudas", 90, body_debts)

        if r.commitments:
            def body_commit(yy0):
                yy = yy0
                for texto, fecha, responsable in r.commitments:
                    for ln in wrap(f"• {texto}", 100):
                        t(left+16, yy, ln, size=9.0, col=muted)
                        yy -= 12
                    t(left+24, yy, f"Fecha: {fecha or '—'}   |   Responsable: {responsable or '—'}", size=8.4, col=accent)
                    yy -= 16
                return yy
            section_card("Compromisos acordados", 90, body_commit)

        draw_footer()

    c.showPage()
    c.save()
    return buf.getvalue()


# ----------------------------
# Word (.docx) — mismo contenido que el TXT/PDF, formato editable
# ----------------------------
def generate_docx(r: ClientFinReport) -> bytes:
    if not DOCX_OK:
        raise RuntimeError("python-docx no está disponible. Agregar `python-docx` a requirements.txt.")

    MUTED = RGBColor(0x55, 0x55, 0x55)
    GOOD = RGBColor(0x1E, 0x8E, 0x5A)
    WARN = RGBColor(0xB8, 0x86, 0x00)
    BAD = RGBColor(0xC0, 0x39, 0x39)
    rating_color = {"good": GOOD, "warn": WARN, "bad": BAD}

    doc = docx.Document()

    doc.add_heading("Diagnóstico de Finanzas Personales — Reporte de Consultoría", level=1)

    p = doc.add_paragraph()
    p.add_run(f"Consultor/a: {r.consultant}   |   Cliente: {r.client}   |   Fecha: {r.report_date}").bold = True
    edad = "—" if r.age is None else str(r.age)
    doc.add_paragraph(f"Edad: {edad}  |  Ocupación: {r.occupation}  |  Dependientes: {r.dependents}")
    doc.add_paragraph(f"Objetivo principal: {r.objective}")

    salud = doc.add_paragraph()
    run = salud.add_run(f"Salud financiera general: {r.health_label} ({r.health_score}/{r.health_max})")
    run.bold = True
    run.font.color.rgb = GOOD if r.health_score >= 6 else (BAD if r.health_score <= 3 else WARN)

    doc.add_heading("Perfil de riesgo", level=2)
    doc.add_paragraph(f"{r.risk_category} — {r.risk_score}/{r.risk_max} pts", style=None).runs[0].bold = True
    doc.add_paragraph(r.risk_description)
    doc.add_paragraph(
        f"Mezcla ilustrativa (no es recomendación de inversión): {r.risk_alloc_fixed}% bajo riesgo/ahorro — "
        f"{r.risk_alloc_variable}% mayor riesgo/crecimiento"
    )
    nota = doc.add_paragraph(RISK_METHOD_NOTE)
    nota.runs[0].italic = True
    nota.runs[0].font.color.rgb = MUTED

    if r.portfolio_tier and r.portfolio_categories:
        doc.add_heading("Cartera modelo sugerida", level=2)
        p_tier = doc.add_paragraph()
        p_tier.add_run(f"Nivel sugerido: {r.portfolio_tier}").bold = True
        if r.portfolio_age_note:
            p_age = doc.add_paragraph(r.portfolio_age_note)
            p_age.runs[0].italic = True
            p_age.runs[0].font.color.rgb = MUTED
        pt = doc.add_table(rows=1, cols=2)
        pt.style = "Light Grid Accent 1"
        hdr = pt.rows[0].cells
        hdr[0].text, hdr[1].text = "Categoría de instrumento", "% sugerido"
        for categoria, pct in r.portfolio_categories:
            row = pt.add_row().cells
            row[0].text = categoria
            row[1].text = f"{pct}%"
        p_disc = doc.add_paragraph(PORTFOLIO_DISCLAIMER)
        p_disc.runs[0].italic = True
        p_disc.runs[0].font.size = Pt(8.5)
        p_disc.runs[0].font.color.rgb = MUTED

    doc.add_heading("Diagnóstico financiero rápido", level=2)
    tabla_datos = [
        ("Ingreso mensual", fmt_pyg(r.income)),
        ("Gastos fijos", fmt_pyg(r.fixed_expenses)),
        ("Gastos variables", fmt_pyg(r.variable_expenses)),
        ("Cuota de deudas", fmt_pyg(r.debt_payment)),
        ("Ahorro/inversión mensual", fmt_pyg(r.savings_monthly)),
        ("Fondo de emergencia actual", fmt_pyg(r.emergency_fund)),
    ]
    if r.net_worth is not None:
        tabla_datos.append(("Patrimonio neto aproximado", fmt_pyg(r.net_worth)))
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, val in tabla_datos:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = val

    doc.add_heading("Indicadores", level=2)
    savings_txt = "No calculable (sin ingreso)" if r.savings_rating.key == "na" else fmt_pct(r.savings_rate)
    dti_txt = "No calculable (sin ingreso)" if r.dti_rating.key == "na" else fmt_pct(r.dti)
    emergency_txt = "No calculable (sin gastos)" if r.emergency_rating.key == "na" else f"{r.emergency_months:.1f} meses"
    available_after_savings = r.cashflow - r.savings_monthly
    indicadores = [
        ("Excedente antes de ahorro", fmt_pyg(r.cashflow), r.cashflow_rating),
        ("Tasa de ahorro", savings_txt, r.savings_rating),
        ("Endeudamiento (cuota/ingreso)", dti_txt, r.dti_rating),
        ("Fondo de emergencia", emergency_txt, r.emergency_rating),
    ]
    for label, val, rating in indicadores:
        para_ind = doc.add_paragraph(style="List Bullet")
        para_ind.add_run(f"{label}: {val} — ").bold = False
        run_rating = para_ind.add_run(f"{rating.label}")
        run_rating.bold = True
        run_rating.font.color.rgb = rating_color.get(rating.key, MUTED)
    p_disp = doc.add_paragraph(style="List Bullet")
    p_disp.add_run(f"Disponible después de ahorro (excedente − ahorro declarado): {fmt_pyg(available_after_savings)}")
    if available_after_savings < 0:
        warn_p = doc.add_paragraph(
            "⚠ El ahorro declarado supera el excedente disponible; verificar con el cliente el origen de "
            "esos fondos."
        )
        warn_p.runs[0].font.color.rgb = BAD

    doc.add_heading("Plan de acción sugerido", level=2)
    for item in r.action_plan:
        doc.add_paragraph(item, style="List Bullet")

    if r.savings_goal_amount is not None and r.savings_goal_months:
        doc.add_heading("Meta de ahorro", level=2)
        nombre_meta = r.savings_goal_name or "Meta sin nombre"
        doc.add_paragraph(f"{nombre_meta}: {fmt_pyg(r.savings_goal_amount)} en {r.savings_goal_months} meses")
        doc.add_paragraph(f"Ahorro mensual necesario: {fmt_pyg(r.savings_goal_required_monthly)}")
        doc.add_paragraph(f"Ahorro mensual actual declarado: {fmt_pyg(r.savings_monthly)}")
        veredicto = "Alcanzable con el ahorro actual" if r.savings_goal_feasible else "Requiere aumentar el ahorro mensual o extender el plazo"
        p_ver = doc.add_paragraph()
        p_ver.add_run(f"Veredicto: {veredicto}").bold = True

    if r.debts:
        doc.add_heading("Inventario de deudas", level=2)
        dt = doc.add_table(rows=1, cols=4)
        dt.style = "Light Grid Accent 1"
        hdr = dt.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Deuda", "Saldo", "Tasa anual", "Pago mínimo"
        for d in r.debts:
            row = dt.add_row().cells
            row[0].text = d.nombre
            row[1].text = fmt_pyg(d.saldo)
            row[2].text = fmt_pct(d.tasa_anual)
            row[3].text = fmt_pyg(d.pago_minimo)
        if r.debt_priority_note:
            doc.add_paragraph(r.debt_priority_note).runs[0].italic = True

    if r.commitments:
        doc.add_heading("Compromisos acordados", level=2)
        ct = doc.add_table(rows=1, cols=3)
        ct.style = "Light Grid Accent 1"
        hdr = ct.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Compromiso", "Fecha", "Responsable"
        for texto, fecha, responsable in r.commitments:
            row = ct.add_row().cells
            row[0].text = texto
            row[1].text = fecha or "—"
            row[2].text = responsable or "—"

    doc.add_paragraph()
    disc = doc.add_paragraph(DISCLAIMER)
    disc.runs[0].italic = True
    disc.runs[0].font.size = Pt(8.5)
    disc.runs[0].font.color.rgb = MUTED

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
